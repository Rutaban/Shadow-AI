"""Minimal Springer LNCS and CCIS layout helpers built on python-docx.

Page geometry and default fonts are copied from the official Springer
proceedings Word template (splnproc2510): A4 paper with a 12.2 by 19.3 cm text
area and Times New Roman throughout.
"""
from __future__ import annotations

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

BODY_PT = 10
SMALL_PT = 9
CAPTION_PT = 9


def _set(el, tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    el.append(node)
    return node


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Twips(11906)
    section.page_height = Twips(16838)
    section.top_margin = Twips(2948)
    section.bottom_margin = Twips(2948)
    section.left_margin = Twips(2494)
    section.right_margin = Twips(2494)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(BODY_PT)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + attr), "Times New Roman")
    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(12)
    return doc


def para(doc, text="", size=BODY_PT, bold=False, italic=False, align="justify",
         before=0, after=0, indent=0.0, line=12, first_line=0.0):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line)
    fmt.left_indent = Cm(indent)
    fmt.first_line_indent = Cm(first_line)
    fmt.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                     "center": WD_ALIGN_PARAGRAPH.CENTER,
                     "left": WD_ALIGN_PARAGRAPH.LEFT,
                     "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    if text:
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
    return p


def rich(doc, chunks, size=BODY_PT, align="justify", before=0, after=0,
         line=12, first_line=0.0, indent=0.0):
    """chunks is a list of (text, bold, italic) triples."""
    p = para(doc, "", size=size, align=align, before=before, after=after,
             line=line, first_line=first_line, indent=indent)
    for chunk in chunks:
        text, bold, italic = chunk[0], chunk[1], chunk[2]
        superscript = chunk[3] if len(chunk) > 3 else False
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.font.superscript = superscript
        run.font.name = "Times New Roman"
    return p


def heading(doc, text, level=1):
    if level == 1:
        return para(doc, text, size=12, bold=True, align="left", before=18, after=6, line=14)
    if level == 2:
        return para(doc, text, size=10, bold=True, align="left", before=12, after=4)
    return para(doc, text, size=10, bold=True, align="left", before=10, after=2)


def body(doc, text, first=False):
    return para(doc, text, align="justify", first_line=0.0 if first else 0.42,
                before=0, after=0)


def figure(doc, path, caption, width_cm=12.0):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    p.paragraph_format.keep_with_next = True
    cap = para(doc, "", size=CAPTION_PT, align="center", after=10, line=11)
    label = cap.add_run(caption.split("|", 1)[0])
    label.bold = True
    label.font.size = Pt(CAPTION_PT)
    label.font.name = "Times New Roman"
    rest = cap.add_run(caption.split("|", 1)[1])
    rest.font.size = Pt(CAPTION_PT)
    rest.font.name = "Times New Roman"
    return cap


def table_caption(doc, caption):
    cap = para(doc, "", size=CAPTION_PT, align="left", before=10, after=3, line=11)
    cap.paragraph_format.keep_with_next = True
    label = cap.add_run(caption.split("|", 1)[0])
    label.bold = True
    label.font.size = Pt(CAPTION_PT)
    label.font.name = "Times New Roman"
    rest = cap.add_run(caption.split("|", 1)[1])
    rest.font.size = Pt(CAPTION_PT)
    rest.font.name = "Times New Roman"
    return cap


def _clear_borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "nil")
        borders.append(node)
    tc_pr.append(borders)


def _rule(cell, edge, size=8):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    existing = borders.find(qn("w:" + edge))
    if existing is not None:
        borders.remove(existing)
    node = OxmlElement("w:" + edge)
    node.set(qn("w:val"), "single")
    node.set(qn("w:sz"), str(size))
    node.set(qn("w:space"), "0")
    node.set(qn("w:color"), "000000")
    borders.append(node)


def table_spacer(doc):
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    sp.paragraph_format.line_spacing = Pt(5)
    return sp


def table(doc, rows, widths=None, size=SMALL_PT, header_rows=1, align_right=None):
    align_right = align_right or []
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.autofit = False
    for r, record in enumerate(rows):
        for c, value in enumerate(record):
            cell = t.cell(r, c)
            _clear_borders(cell)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(size + 1.6)
            if c in align_right and r >= header_rows:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(value))
            run.font.size = Pt(size)
            run.font.name = "Times New Roman"
            run.bold = r < header_rows
            if widths:
                cell.width = Cm(widths[c])
    for c in range(len(rows[0])):
        _rule(t.cell(0, c), "top", 8)
        _rule(t.cell(header_rows - 1, c), "bottom", 6)
        _rule(t.cell(len(rows) - 1, c), "bottom", 8)
    table_spacer(doc)
    return t


def references(doc, entries):
    for index, entry in enumerate(entries, start=1):
        p = para(doc, "", size=SMALL_PT, align="justify", line=11,
                 indent=0.62, after=1)
        p.paragraph_format.first_line_indent = Cm(-0.62)
        run = p.add_run(f"{index}. {entry}")
        run.font.size = Pt(SMALL_PT)
        run.font.name = "Times New Roman"
